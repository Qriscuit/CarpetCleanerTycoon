from pathlib import Path
import math, re, json
import xlsxwriter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'design'
XLSX = ROOT / 'Carpet_Cleaner_Economy_v0_2.xlsx'
DOCX = ROOT / 'Carpet_Cleaner_Tycoon_GDD_v0_2.docx'

# Numeric defaults live here; workbook input cells are editable for playtesting.
shops = [
 ['S1','Neighborhood Shop',0,20,50,10,40,30,0,10,0,0,'Brush; mint/cream tiles','Playable slice'],
 ['S2','Busy High Street',600,60,60,10,40,30,36,25,15,5,'Jet + squeegee; coral storefront','Playable slice'],
 ['S3','Restoration Studio',6000,150,100,10,36,24,30,65,12,20,'Treatment + pile kit; indigo/timber','Placeholder'],
 ['S4','Commercial Workshop',40000,350,130,10,36,24,30,175,15,30,'Broad tools; teal industrial bay','Placeholder'],
 ['S5','Flagship Atelier',180000,900,180,10,30,20,24,450,18,40,'Integrated rig; cream/gold atelier','Placeholder'],
]
rates = [min(x[6],x[7],x[8] or x[7])*x[9] for x in shops]
orders = [min(x[6],x[7],x[8] or x[7]) for x in shops]
cycles = [x[4]+x[5] for x in shops]
wb = xlsxwriter.Workbook(XLSX)
wb.set_properties({'title':'Carpet Cleaner — economy v0.2','subject':'Brush-first metagame balance','author':'Carpet Cleaner design','comments':'Proposed tuning. First two shops only are the playable slice.'})
wb.set_calc_mode('auto')
navy,teal,cream,blue = '#203B43','#167E78','#F5F2E9','#1868BB'
fm = {}
fm['title']=wb.add_format({'bold':True,'font_size':22,'font_color':'white','bg_color':navy,'valign':'vcenter'})
fm['sub']=wb.add_format({'font_color':'#52666C','text_wrap':True,'valign':'vcenter'})
fm['head']=wb.add_format({'bold':True,'font_color':'white','bg_color':teal,'text_wrap':True,'valign':'vcenter','border':0})
for kind,color,bg in [('txt',navy,'white'),('input',blue,'#EEF5FF'),('calc',navy,'#EAF5F1')]:
 for typ,fmt in [('','General'),('n','#,##0.0;[Red](#,##0.0);–'),('i','#,##0;[Red](#,##0);–'),('p','0%')]:
  fm[kind+typ]=wb.add_format({'font_color':color,'bg_color':bg,'text_wrap':kind=='txt','valign':'top','num_format':fmt,'bottom':1,'bottom_color':'#E1E8E6'})
fm['link']=wb.add_format({'font_color':teal,'underline':True,'text_wrap':True})
fm['good']=wb.add_format({'bg_color':'#CDEDDD','font_color':'#205D45'})
fm['warn']=wb.add_format({'bg_color':'#FFE3C0','font_color':'#874C11'})

def sheet(name,title,subtitle,headers,widths):
 s=wb.add_worksheet(name); s.hide_gridlines(2); s.set_tab_color(teal)
 n=len(headers); s.merge_range(0,0,0,n-1,title,fm['title']);s.set_row(0,38)
 s.merge_range(1,0,2,n-1,subtitle,fm['sub']);s.set_row(1,24);s.set_row(2,22)
 s.write_row(4,0,headers,fm['head']);s.set_row(4,34)
 for c,w in enumerate(widths):s.set_column(c,c,w)
 s.freeze_panes(5,1);s.set_landscape();s.set_paper(9);s.fit_to_pages(1,0);s.repeat_rows(0,4)
 s.set_margins(.3,.3,.4,.4);s.set_footer('&LCarpet Cleaner • proposed tuning&RPage &P of &N')
 return s
def val(s,r,c,v,kind='txt',typ=''):
 s.write(r,c,v,fm[kind+typ])
def formula(s,r,c,f,v,typ='n'):
 s.write_formula(r,c,f,fm['calc'+typ],v)
def numeric(s,first,last,col,minimum=0,maximum=100000000,integer=False):
 s.data_validation(first,col,last,col,{'validate':'integer' if integer else 'decimal','criteria':'between','minimum':minimum,'maximum':maximum,'input_title':'Tuning input','input_message':'Change blue cells; green cells calculate.','error_title':'Check this value','error_message':f'Enter a value between {minimum} and {maximum}.'})

guide=sheet('Read Me','BRUSH → BONZI → A GROWING CHAIN','Economy v0.2 | 11 September 2026 | Blue = editable input • Green = formula • All timings are hypotheses, not observed playtests.',['Sheet / topic','Purpose','How to use'],[27,66,76])
guide_rows=[
 ('Assumptions','Shared economy and scenario inputs.','Change these and blue Shops/Upgrades values. Formula caches are supplied for previews; Excel/LibreOffice recalculates edits.'),
 ('Shops','Five locations: prices, job rewards, service capacities, local gates.','Shops 1–2 are the slice. Shops 3–5 are placeholders. Zero finishing means the service is not required.'),
 ('Items','Personal tools and blueprint/build ownership.','Whole blueprints are permanent. Starter and opening kits have no hidden construction cost.'),
 ('Upgrades','Optional local machine purchases with bottleneck-aware output and payback.','Named before/after scenarios are linked sequentially. Payback is cash-only; parts/gates require extra effort.'),
 ('Pacing','Sequential progression with cash carried across purchases.','Average-rate model; active job rounding, local manual/auto gates. No optional purchases or offline cap in this sheet.'),
 ('Starter Routes','Three explicit ways to reach High Street.','Continuous active, 50% activity and minimum manual work followed by absence. Remainder cash retained.'),
 ('Offline','Eight-hour cap, integer completed deliveries, owned/enabled branches.','Change hours, ownership switches and fractional order progress. S1 enabled means Bonzi has been built.'),
 ('Blueprints & Parts','Guaranteed unlocks and optional crafting costs.','No fragment currency, random drop requirement, parts from idle, or mandatory parts gate.'),
 ('UI & States','UI states and transitions for the coming implementation.','Design only; no runtime/UI changes are included.'),
 ('Checks','Formula checks and default-route simulation results.','Formula checks update with workbook inputs. The discrete route check is a labeled default snapshot.'),
 ('Scope','Old shops keep earning; first shop free; parts begin in S2.','No energy, rent, gems, chests, prestige, monetization or operating expenses in this model.'),
 ('Model limits','Continuous pacing approximates discrete automatic deliveries.','Offline uses whole orders. Pacing assumes eligible certification is within the required manual count, immediate claims, base kit and constant average activity.'),
 ('Source of truth','GDD v0.2 gives rules; workbook gives numeric inputs.','Written examples match the initial defaults. If workbook inputs change, update written examples before implementation.'),
]
for r,row in enumerate(guide_rows,5):
 guide.set_row(r,46)
 for c,v in enumerate(row):val(guide,r,c,v)
 if row[0] in ['Assumptions','Shops','Items','Upgrades','Pacing','Starter Routes','Offline','Blueprints & Parts','UI & States','Checks']:
  guide.write_url(r,0,f"internal:'{row[0]}'!A1",fm['link'],row[0])

assum=sheet('Assumptions','THE TUNING DESK','Input values apply to the proposed slice. No real-money values. Active fraction describes time spent completing manual jobs while production runs.',['Parameter','Value','Unit / meaning','Status'],[35,19,82,25])
assum_rows=[
 ('Starting cash',0,'cash; starter premises and brush already owned','Proposed'),
 ('Bonzi build cost',100,'cash; includes first automated chassis','Proposed'),
 ('Bonzi blueprint at job',3,'local S1 paid manual completions','Proposed'),
 ('Offline cap',8,'hours per absence, not a live production cap','Proposed'),
 ('Active fraction',1,'Pacing scenario; >0 to 1. 0 is invalid because manual gates exist.','Scenario'),
 ('Clump completion',.9,'fraction cleared before Finish is allowed; not current toy rule','Proposed'),
 ('Surface completion',.9,'fraction of surface cleaned; must also pass clump check','Proposed'),
 ('Parts per eligible manual',1,'integer; S2+ rewarded jobs only','Proposed'),
 ('Wide Brush build cost',80,'cash; blueprint at eight S1 manual jobs','Proposed'),
 ('Wide Brush work seconds',40,'seconds on the same S1 rug; handling stays unchanged','Unmeasured target'),
 ('Slow activity scenario',.5,'Starter Routes only; after Bonzi','Scenario'),
]
for r,row in enumerate(assum_rows,5):
 assum.set_row(r,32)
 for c,v in enumerate(row):val(assum,r,c,v,'input' if c==1 else 'txt','p' if c==1 and r in [9,10,11,15] else 'n' if c==1 else '')
 numeric(assum,r,r,1,0,100000,False)
for r in [9,15]:numeric(assum,r,r,1,.01,1)
for r in [10,11]:numeric(assum,r,r,1,.5,1)
numeric(assum,12,12,1,0,1000,True)
numeric(assum,7,7,1,1,10000,True)

ss=sheet('Shops','THE FIVE-SHOP ARC','Opening price includes the tools and local automated unit needed to operate. S1 automation is separately enabled by Bonzi. All rates below assume automation enabled.',['ID','Shop','Open cash','Manual cash','Work sec','Handling sec','Demand /h','Cleaner /h','Finish /h (0=N/A)','Routine cash/order','Prev-shop manual gate','Prev-shop auto gate','Identity / opening kit','Scope','Cycle sec','Auto orders/h','Auto cash/h','Manual cash/h','Chain auto cash/h'],[9,25,14,14,12,12,13,13,16,16,17,16,42,18,12,14,14,16,18])
for r,x in enumerate(shops,5):
 ss.set_row(r,60);i=r-5;e=r+1
 for c,v in enumerate(x):val(ss,r,c,v,'input' if 2<=c<=11 else 'txt','i' if 2<=c<=11 else '')
 formula(ss,r,14,f'=E{e}+F{e}',cycles[i])
 formula(ss,r,15,f'=MIN(G{e},H{e},IF(I{e}=0,H{e},I{e}))',orders[i])
 formula(ss,r,16,f'=P{e}*J{e}',rates[i])
 formula(ss,r,17,f'=D{e}*3600/O{e}',x[3]*3600/cycles[i])
 formula(ss,r,18,f'=SUM(Q$6:Q{e})',sum(rates[:i+1]))
for c in range(2,12):numeric(ss,5,9,c,1 if c in [3,4,6,7,9] else 0,10000000,c in [10,11])
ss.autofilter(4,0,9,18)
ss.conditional_format('Q6:Q10',{'type':'data_bar','bar_color':teal})

its=sheet('Items','KNOW IT → BUILD IT → USE IT','No duplicate personal tool builds. Opening kits are granted as part of the shop purchase; “included” never means a second bill. Later optional variant prices remain unset.',['Item','Unlock / blueprint','Build cash','Parts','Ownership','Effect / boundary'],[27,50,17,13,28,77])
item_rows=[
 ['Hand Brush','First launch',0,0,'Account personal tool','Only cleaning tool in the starter economy. Existing prototype display tools stay prototype-only until wet service.'],
 ['Bonzi + basic chassis','3 paid S1 manual jobs',100,0,'Named buddy + S1 unit','30 orders/h capacity. A new unit is bundled with each later shop.'],
 ['Wide Brush','8 paid S1 manual jobs',80,0,'Account personal tool','S1 work target 50→40 seconds; no handling change. Optional.'],
 ['Jet + Hand Squeegee','High Street opening','Included',0,'Account personal tools','Required wet-service capability. Recipe must function before S2 is playable.'],
 ['Restoration starter kit','Restoration Studio opening','Included',0,'Account personal tools','Treatment and pile capability; exact tool interaction is later scope.'],
 ['Broad cleaning kit','Commercial Workshop opening','Included',0,'Account personal tools','Wide work surface capability; later scope.'],
 ['Integrated Rig','Flagship opening','Included',0,'Account personal tool','Capstone starter capability; later scope.'],
 ['Foam brush / rake / CRB / rotary variants','Later authored milestones','TBD','TBD','Personal optional variants','Do not price capabilities that are not yet designed. No random required unlocks.'],
 ['Local modules','See Upgrades and Blueprints & Parts','See Upgrades','See Upgrades','Bound to purchased branch','Incremental one-time named upgrades. No resale or infinite levels.'],
]
for r,row in enumerate(item_rows,5):
 its.set_row(r,48)
 for c,v in enumerate(row):val(its,r,c,v)
formula(its,6,2,'=Assumptions!B7',100,'i');formula(its,7,2,'=Assumptions!B14',80,'i')

up=sheet('Upgrades','SPEND ON THE ACTUAL BOTTLENECK','Rows are named scenarios, not purchases to sum together. Sequential rows inherit the previous output settings. Payback ignores parts acquisition time. Zero gain is shown explicitly.',['Scenario / purchase','Cash','Parts','Before demand','Before cleaner','Before finish (0=N/A)','After demand','After cleaner','After finish (0=N/A)','Cash/order','Before orders/h','After orders/h','Before cash/h','After cash/h','Gain cash/h','Cash payback h','Prerequisite / meaning'],[28,13,11,14,14,17,14,14,17,13,14,14,14,14,14,16,60])
urows=[
 ['S1 Mk II only',120,0,40,30,0,40,45,0,10,'Bonzi + 10 S1 manual jobs; demand becomes bottleneck'],
 ['S1 intake after Mk II',100,0,40,45,0,60,45,0,10,'Owned Mk II; cleaner becomes bottleneck'],
 ['S1 intake BEFORE Mk II',100,0,40,30,0,60,30,0,10,'Alternative order: no immediate income gain'],
 ['S2 cleaner module',900,6,40,30,36,40,45,36,25,'6 local manual jobs; finishing becomes bottleneck'],
 ['S2 finish after cleaner',600,4,40,45,36,40,45,48,25,'Owned cleaner module; demand becomes bottleneck'],
 ['S2 intake after both',750,0,40,45,48,60,45,48,25,'Owned cleaner + finish; cleaner becomes bottleneck'],
]
for r,x in enumerate(urows,5):
 e=r+1;up.set_row(r,52)
 for c,v in enumerate(x[:10]):val(up,r,c,v,'input' if c in [1,2,6,7,8] else 'txt','n' if 1<=c<=9 else '')
 val(up,r,16,x[10])
 # Before settings reference base shop or the preceding sequential scenario.
 base=6 if r<=7 else 7
 previous={6:6,9:9,10:10}.get(r)
 for c,source in [(3,'G'),(4,'H'),(5,'I')]:
  f=f'={source}{previous}' if previous else f'=Shops!{source}{base}'
  formula(up,r,c,f,x[c])
 formula(up,r,9,f'=Shops!J{base}',x[9])
 # Unchanged after capacities link to before, leaving only the changed setting editable.
 changed={5:7,6:6,7:6,8:7,9:8,10:6}[r]
 for c in [6,7,8]:
  if c!=changed:formula(up,r,c,f'={chr(65+c-3)}{e}',x[c])
 b=min(x[3],x[4],x[5] or x[4]);a=min(x[6],x[7],x[8] or x[7]);gain=(a-b)*x[9]
 for c,f,v in [(10,f'=MIN(D{e},E{e},IF(F{e}=0,E{e},F{e}))',b),(11,f'=MIN(G{e},H{e},IF(I{e}=0,H{e},I{e}))',a),(12,f'=K{e}*J{e}',b*x[9]),(13,f'=L{e}*J{e}',a*x[9]),(14,f'=N{e}-M{e}',gain),(15,f'=IF(O{e}>0,B{e}/O{e},"No gain")',x[1]/gain if gain>0 else 'No gain')]:formula(up,r,c,f,v)
 for c in [1,2,changed]:numeric(up,r,r,c,0,100000, c==2)
up.conditional_format('O6:O11',{'type':'cell','criteria':'<=','value':0,'format':fm['warn']})
up.write_row(14,0,['Personal upgrade','Cash','Base manual /h','Upgraded manual /h','Gain /h','Active payback h'],fm['head'])
val(up,15,0,'Wide Brush')
formula(up,15,1,'=Assumptions!B14',80)
formula(up,15,2,'=Shops!R6',1200)
formula(up,15,3,'=Shops!D6*3600/(Assumptions!B15+Shops!F6)',1440)
formula(up,15,4,'=D16-C16',240)
formula(up,15,5,'=IF(E16>0,B16/E16,"No gain")',1/3)

# Sequential average-rate pacing, with full manual job rounding and carry-over cash.
pace=sheet('Pacing','HOW LONG UNTIL THE NEXT DOOR?','Constant online activity scenario from Assumptions. Whole manual jobs, average bot income, base kits; no optional spending. Certification counts inside the local manual minimum. Later-shop outputs are placeholders.',['Milestone','Cost','Cash entering','Manual cash/job','Cycle sec','Active share','Manual cash/h','Prior-chain auto cash/h','Manual cash still needed?','Manual jobs gate left','Auto deliveries gate','Local auto orders/h','Hours needed (continuous)','Manual jobs done','Elapsed h rounded','Manual income','Auto income','Cash after purchase','Total elapsed h'],[24,13,14,15,12,13,16,18,19,18,17,18,19,17,18,16,16,18,18])
pacing=[]
cash=0;elapsed=0
for i in range(5):
 r=5+i;e=r+1;srow=5+i # source shop Excel row: S1 for second milestone is 6
 if i==0:
  cost=100;reward=20;cycle=60;auto=0;mgate=3;agate=0;local=0
 else:
  x=shops[i-1];cost=shops[i][2];reward=x[3];cycle=cycles[i-1];auto=sum(rates[:i]);mgate=shops[i][10]-(pacing[0]['jobs'] if i==1 else 0);mgate=max(0,mgate);agate=shops[i][11];local=orders[i-1]
 active=1;mh=reward*3600/cycle;need=max(0,cost-cash)
 raw=max(need/(mh+auto),mgate*cycle/3600,agate/local if local else 0)
 jobs=math.ceil(raw*3600/cycle-1e-9);hours=max(jobs*cycle/3600,agate/local if local else 0)
 mi=jobs*reward;ai=hours*auto;after=cash+mi+ai-cost;elapsed+=hours
 pacing.append({'jobs':jobs,'hours':hours,'cash':after,'elapsed':elapsed})
 vals=['Build Bonzi' if i==0 else 'Open '+shops[i][1],cost,cash,reward,cycle,active,mh,auto,need,mgate,agate,local,raw,jobs,hours,mi,ai,after,elapsed]
 pace.set_row(r,48);val(pace,r,0,vals[0])
 expr={1:'=Assumptions!B7' if i==0 else f'=Shops!C{i+6}',2:'=Assumptions!B6' if i==0 else f'=R{e-1}',3:'=Shops!D6' if i==0 else f'=Shops!D{srow}',4:'=Shops!O6' if i==0 else f'=Shops!O{srow}',5:'=Assumptions!B10',6:f'=D{e}*3600/E{e}*F{e}',7:'=0' if i==0 else f'=SUM(Shops!Q$6:Q{srow})',8:f'=MAX(0,B{e}-C{e})',9:'=Assumptions!B8' if i==0 else f'=MAX(0,Shops!K{i+6}'+('-N6)' if i==1 else ')'),10:'=0' if i==0 else f'=Shops!L{i+6}',11:'=0' if i==0 else f'=Shops!P{srow}',12:f'=MAX(I{e}/(G{e}+H{e}),J{e}*E{e}/(3600*F{e}),IF(K{e}=0,0,K{e}/MAX(L{e},0.000001)))',13:f'=ROUNDUP(ROUND(M{e}*3600*F{e}/E{e},9),0)',14:f'=MAX(N{e}*E{e}/(3600*F{e}),IF(K{e}=0,0,K{e}/MAX(L{e},0.000001)))',15:f'=N{e}*D{e}',16:f'=O{e}*H{e}',17:f'=C{e}+P{e}+Q{e}-B{e}',18:f'=SUM(O$6:O{e})'}
 for c,f in expr.items():formula(pace,r,c,f,vals[c],'p' if c==5 else 'i' if c in [9,10,13] else 'n')
 cash=after
chart=wb.add_chart({'type':'column'})
chart.add_series({'name':'Hours from previous milestone','categories':'=Pacing!$A$6:$A$10','values':'=Pacing!$O$6:$O$10','fill':{'color':teal},'border':{'none':True}})
chart.set_title({'name':'Later shops deliberately take longer'});chart.set_y_axis({'name':'Elapsed online hours'});chart.set_legend({'none':True});chart.set_size({'width':820,'height':330});pace.insert_chart('A14',chart)

sr=sheet('Starter Routes','THREE WAYS TO REACH HIGH STREET','Routes start at launch, buy Bonzi first, skip optional upgrades. Guided return does only the required manual jobs and then waits. Continuous route estimates use average routine income.',['Route','Active share after Bonzi','Bonzi minute','Post-Bonzi jobs','Post-Bonzi active/online min','Away min','Total elapsed min','Cash remaining','Interpretation'],[28,20,17,18,25,17,22,18,75])
for r,frac in [(5,1),(6,.5),(7,1)]:
 e=r+1;sr.set_row(r,58)
 val(sr,r,0,['Keep cleaning','Half-time activity','Minimum work, then return'][r-5])
 formula(sr,r,1,'=1' if r!=6 else '=Assumptions!B16',frac,'p')
 formula(sr,r,2,'=Pacing!N6*Shops!O6/60',5)
 if r<7:
  needh=max(600/(1200*frac+300),10/60/frac,5/30)
  j=math.ceil(needh*60*frac-1e-9);mins=j/frac
  formula(sr,r,3,f'=ROUNDUP(ROUND(MAX(MAX(0,Shops!C7-Pacing!R6)/(Shops!R6*B{e}+Shops!Q6),MAX(0,Shops!K7-Pacing!N6)*Shops!O6/(3600*B{e}),Shops!L7/Shops!P6)*3600*B{e}/Shops!O6,9),0)',j,'i')
  formula(sr,r,4,f'=D{e}*Shops!O6/(60*B{e})',mins)
  formula(sr,r,5,'=0',0)
  rem=j*20+mins*5-600
 else:
  j=10;mins=10;rem=0
  formula(sr,r,3,'=MAX(0,Shops!K7-Pacing!N6)',10,'i')
  formula(sr,r,4,f'=D{e}*Shops!O6/60',10)
  formula(sr,r,5,f'=MAX(0,(Shops!C7-Pacing!R6-D{e}*Shops!D6-E{e}*Shops!Q6/60)/(Shops!Q6/60),Shops!L7/Shops!P6*60-E{e})',70)
 formula(sr,r,6,f'=C{e}+E{e}+F{e}',29 if r==5 else 45 if r==6 else 85)
 formula(sr,r,7,f'=Pacing!R6+D{e}*Shops!D6+(E{e}+F{e})*Shops!Q6/60-Shops!C7',rem)
 val(sr,r,8,['29 jobs total; 12 bot deliveries after purchase at defaults.','25 jobs total; slower manual activity, more routine income.','15 jobs total. If away time exceeds cap, multiple return intervals are needed.'][r-5])

off=sheet('Offline','WHAT HAPPENS WHILE YOU ARE AWAY','Whole automatic orders pay; fractional order progress carries forward. Ownership switch means the branch already existed and its automation was enabled at departure. No upgrade/new shop is applied retrospectively.',['Shop','Owned + enabled (0/1)','Away hours','Credited hours','Auto orders/h','Cash/order','Order remainder entering','Completed orders','Cash earned','Order remainder leaving'],[28,23,16,18,19,17,24,20,18,24])
for r,x in enumerate(shops,5):
 e=r+1;i=r-5;enabled=1 if i<2 else 0;off.set_row(r,34)
 val(off,r,0,x[1]);val(off,r,1,enabled,'input','i');val(off,r,2,8,'input','n');val(off,r,6,0,'input','n')
 formula(off,r,3,f'=IF(B{e}=1,MIN(MAX(0,C{e}),Assumptions!B9),0)',8*enabled)
 formula(off,r,4,f'=Shops!P{e}',orders[i]);formula(off,r,5,f'=Shops!J{e}',x[9])
 formula(off,r,7,f'=IF(B{e}=1,ROUNDDOWN(ROUND(D{e}*E{e}+G{e},9),0),0)',orders[i]*8*enabled,'i')
 formula(off,r,8,f'=H{e}*F{e}',rates[i]*8*enabled,'i')
 formula(off,r,9,f'=IF(B{e}=1,ROUND(D{e}*E{e}+G{e}-H{e},9),G{e})',0)
 numeric(off,r,r,1,0,1,True);numeric(off,r,r,2,0,100000);numeric(off,r,r,6,0,.999999)
val(off,11,0,'TOTAL CLAIM');formula(off,11,8,'=SUM(I6:I10)',8400,'i')
val(off,13,0,'Cap behavior');off.merge_range(13,1,14,9,'Cap truncates simulated absence time, not the wallet. New store clocks start on purchase. Negative elapsed time earns zero. Reopening cannot award a previously settled interval again.',fm['txt'])

bp=sheet('Blueprints & Parts','PERMANENT KNOWLEDGE, USEFUL PARTS','No research fragments. All required blueprints and capabilities are guaranteed. Parts begin in High Street and are optional for expansion.',['Design / rule','Unlock requirement','Cash cost','Parts cost','Consumption / result'],[32,57,17,17,76])
bprows=[
 ['Bonzi','3 paid S1 manual jobs',100,0,'Whole blueprint permanent; cash builds companion + local unit.'],
 ['Wide Brush','8 paid S1 manual jobs',80,0,'Whole blueprint permanent; optional personal build.'],
 ['S1 Mk II','10 paid S1 manual jobs; Bonzi built',120,0,'Local module; first bot improvement teaches cash-only upgrades.'],
 ['S1 intake','10 paid S1 manual jobs; Bonzi built',100,0,'Local capacity upgrade; can be bought before Mk II but preview zero gain.'],
 ['S2 starter kit','Buy S2','Included',0,'Jet, squeegee, basic local unit and finishing station usable immediately.'],
 ['S2 cleaner','6 paid S2 manual jobs',900,6,'Optional; subtract six parts on successful build only.'],
 ['S2 finishing','6 paid S2 manual jobs',600,4,'Optional; may be bought before cleaner; Upgrades shows a suggested sequence.'],
 ['S2 intake','6 paid S2 manual jobs',750,0,'Optional; no parts consumed.'],
 ['Parts source','Each paid manual job in S2+','N/A',1,'Integer award. No parts from idle, S1 or free replay.'],
 ['Duplicate blueprint','Repeated milestone/load','N/A',0,'No-op; grant ID prevents duplication. No conversion currency.'],
 ['Later mandatory kits','Respective shop purchase','Included',0,'Necessary capability never depends on an optional module or random drop.'],
]
for r,row in enumerate(bprows,5):
 bp.set_row(r,43)
 for c,v in enumerate(row):val(bp,r,c,v)
for r,src in [(5,6),(6,7)]:formula(bp,r,2,f'=Items!C{src+1}',100 if r==5 else 80,'i')
for r,ur in [(7,6),(8,7),(10,9),(11,10),(12,11)]:
 formula(bp,r,2,f'=Upgrades!B{ur}',urows[ur-6][1],'i');formula(bp,r,3,f'=Upgrades!C{ur}',urows[ur-6][2],'i')
formula(bp,13,3,'=Assumptions!B13',1,'i')
bp.write_row(18,0,['Parts budget example','Eligible S2+ jobs','Parts earned','Parts spent','Parts remaining'],fm['head'])
val(bp,19,0,'Build S2 cleaner');val(bp,19,1,6,'input','i');formula(bp,19,2,'=B20*Assumptions!B13',6,'i');formula(bp,19,3,'=Upgrades!C9',6,'i');formula(bp,19,4,'=C20-D20',0,'i')
val(bp,20,0,'Then finishing');val(bp,20,1,4,'input','i');formula(bp,20,2,'=B21*Assumptions!B13',4,'i');formula(bp,20,3,'=Upgrades!C10',4,'i');formula(bp,20,4,'=E20+C21-D21',0,'i')
numeric(bp,19,20,1,0,10000,True)
bp.conditional_format('E20:E21',{'type':'cell','criteria':'<','value':0,'format':fm['warn']})

ui=sheet('UI & States','THE NEXT UI PASS','This is a design specification, not an implemented game UI. Expose only the next meaningful decision.',['Surface / system','States','Action','Guard / saved outcome'],[28,56,43,89])
uir=[
 ['Contract','Offered → in progress → finish-ready → paid','Start / resume / Finish','Both cleaning thresholds required. Persist masks/job ID. One payout and local count per job.'],
 ['Bonzi card','Hidden → blueprint known → affordable → built','Build Bonzi','Blueprint and cash checked; debit once; start automation at purchase timestamp.'],
 ['Upgrade card','Blueprint locked / known / affordable / owned','Build or upgrade','Show exact before/after gain and limiting station. Debit cash/parts atomically.'],
 ['Shop expansion','Hidden → revealed → eligible → affordable → owned','Open shop','Local service counts and certification, Bonzi for S2, and cash. One summary includes kit and remaining cash.'],
 ['Owned-shop selector','Current / owned inactive','Visit shop','Free view change. Keep old branches producing. Preserve one chain-wide active personal job.'],
 ['Active job during move','In progress / paused','Finish or explicitly pause','No second manual job until resumed job completes or is abandoned. No abandonment reward/reroll.'],
 ['Return summary','No earnings / claimable / capped','Collect once','Settle past interval once, preserving fractional orders. Offline is not extra income over online.'],
 ['Parts display','Hidden until S2 / visible','View module costs','Only one parts type. No inventory cap, expiry, passive source or mandatory use.'],
 ['Purchase recovery','Ready / saving / committed / failed','Single transaction','Validate again at commit; prevent negative wallet and repeat taps; no duplicate kit or charge after reload.'],
 ['Production','Not enabled / demand limited / cleaner limited / finish limited','Inspect bottleneck','Separate manual and routine pools. Settle old rate before module or price changes.'],
]
for r,row in enumerate(uir,5):
 ui.set_row(r,62)
 for c,v in enumerate(row):val(ui,r,c,v)

# Independent discrete starter simulation: manual completions and bot deliveries award integer cash.
cash=0;manual=0;bot=0;built=False;bonzi_t=None;expand_t=None
for second in range(1,7201):
 if built and (second-bonzi_t)%120==0:cash+=10;bot+=1
 if second%60==0:cash+=20;manual+=1
 if not built and manual>=3 and cash>=100:cash-=100;built=True;bonzi_t=second
 if built and manual>=15 and bot>=5 and cash>=600:
  cash-=600;expand_t=second;break
assert (bonzi_t,expand_t,manual,bot,cash)==(300,1740,29,12,0)
checks=sheet('Checks','MODEL CHECKS & PLAYTEST HANDOFF','Formula checks update when inputs recalculate. Discrete simulation below is a default-only verification, not a dynamic workbook simulation. Test game save/reward behavior separately when implemented.',['Check','Actual','Expected / condition','Result'],[54,25,50,20])
cr=[
 ['Starter premises have no price','=Shops!C6',0,'=0'],
 ['Bonzi first-route cash is nonnegative','=Pacing!R6',0,'>=0'],
 ['S2 route cash is nonnegative','=Pacing!R7',0,'>=0'],
 ['S1 intake-first output gain','=Upgrades!O8',0,'=0'],
 ['All modeled purchase wallets stay nonnegative','=MIN(Pacing!R6:R10)',0,'>=0'],
 ['No offline remainder outside [0,1)','=IF(AND(MIN(Offline!J6:J10)>=0,MAX(Offline!J6:J10)<1),1,0)',1,'=1'],
 ['Guided return fits one offline interval','=IF(\'Starter Routes\'!F8<=Assumptions!B9*60,1,0)',1,'=1'],
 ['Parts example never overspends','=MIN(\'Blueprints & Parts\'!E20:E21)',0,'>=0'],
 ['Bonzi route meets blueprint gate','=Pacing!N6-Assumptions!B8',2,'>=0'],
 ['Opening paid-job count is sufficient','=Pacing!N6+Pacing!N7-Shops!K7',14,'>=0'],
]
for r,(label,f,v,cond) in enumerate(cr,5):
 e=r+1;checks.set_row(r,34);val(checks,r,0,label);formula(checks,r,1,f,v);val(checks,r,2,cond)
 formula(checks,r,3,f'=IF(B{e}{cond},"PASS","REVIEW")','PASS','')
checks.conditional_format('D6:D15',{'type':'text','criteria':'containing','value':'PASS','format':fm['good']})
checks.conditional_format('D6:D15',{'type':'text','criteria':'containing','value':'REVIEW','format':fm['warn']})
val(checks,17,0,'Default discrete event simulation')
for r,row in enumerate([['Bonzi purchase minute',bonzi_t/60,'100 cash after 5 full jobs','PASS'],['High Street opening minute',expand_t/60,'29 jobs, 12 bot deliveries, 0 cash left','PASS'],['Model sensitivity','Change work sec 50→110','Bonzi becomes 10 min; replay route, do not retain timing claims','PLAYTEST'],['Runtime tests needed','Not implemented','Repeat Finish, double-buy, resume, old-shop income, negative clock, offline cap','TODO']],18):
 checks.set_row(r,44)
 for c,v in enumerate(row):val(checks,r,c,v)
for ws in wb.worksheets():ws.set_zoom(85)
wb.close()

# Produce an editable Word GDD from the companion markdown, with proper heading/table styles.
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string('203B43');normal.paragraph_format.space_after=Pt(7)
for h,size in [('Title',31),('Heading 1',19),('Heading 2',13)]:
 st=doc.styles[h];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('167E78');st.paragraph_format.space_before=Pt(16);st.paragraph_format.space_after=Pt(7)
header=sec.header.paragraphs[0];header.text='CARPET CLEANER  /  DESIGN v0.2';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.text='Working proposal • 11 September 2026                                               '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
lines=(OUT/'Metagame_v0_2.md').read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   cells=[x.strip() for x in lines[i].strip('|').split('|')]
   if not all(re.fullmatch(r'[-: ]+',c) for c in cells):rows.append(cells)
   i+=1
  table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
  for j,v in enumerate(rows[0]):table.rows[0].cells[j].text=v
  repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
  for row in rows[1:]:
   cells=table.add_row().cells
   for j,v in enumerate(row):cells[j].text=v
  for rr in table.rows:
   trPr=rr._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
   for cell in rr.cells:
    for p in cell.paragraphs:
     p.paragraph_format.space_after=Pt(5)
     for run in p.runs:run.font.size=Pt(9)
  doc.add_paragraph();continue
 if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
 elif line.startswith('## '):
  p=doc.add_paragraph(line[3:],style='Heading 1')
  if line.startswith('## 1. '):p.paragraph_format.page_break_before=True
 elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
 elif re.match(r'^\d+\. ',line):doc.add_paragraph(re.sub(r'^\d+\. ','',line),style='List Number')
 else:doc.add_paragraph(line)
 i+=1
doc.core_properties.title='Carpet Cleaner Tycoon — GDD v0.2';doc.core_properties.subject='Brush-first economy and shop progression';doc.core_properties.author='Carpet Cleaner design'
doc.save(DOCX)
(OUT/'balance_validation.json').write_text(json.dumps({'default_discrete_first_route':{'bonzi_min':bonzi_t/60,'high_street_min':expand_t/60,'manual_jobs':manual,'bot_deliveries':bot,'cash_remaining':cash},'base_auto_rates':rates,'pacing_baseline':pacing},indent=2),encoding='utf-8')
print(json.dumps({'xlsx':str(XLSX),'gdd':str(DOCX),'pacing':pacing,'discrete_first_expansion_min':expand_t/60},indent=2))
